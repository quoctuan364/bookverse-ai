"""Tạo DOCX báo cáo BookVerse từ Markdown và áp dụng style cố định.

File nguồn Markdown vẫn là nguồn nội dung dễ review; DOCX là artifact bàn giao.
"""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ACADEMIC_REPORT_DRAFT.md"
OUTPUT = ROOT / "outputs" / "Bao_cao_do_an_Luong_Nguyen_Quoc_Tuan_22050098.docx"
ASSET_DIR = ROOT / "outputs" / "report_assets"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(92, 105, 117)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
WHITE = "FFFFFF"


def set_font(run, *, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.extend([begin, instr, separate, text, end])


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor : match.start()]), size=11)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=11, bold=True, color=NAVY)
        else:
            set_font(
                paragraph.add_run(token[1:-1]),
                name="Consolas",
                size=9.5,
                color=DARK_BLUE,
            )
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]), size=11)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_running_furniture(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("BOOKVERSE AI  |  BÁO CÁO ĐỒ ÁN TỐT NGHIỆP")
    set_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, size=8.5, color=MUTED)
    add_field(paragraph, "PAGE", "1")


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    university = doc.add_paragraph()
    university.alignment = WD_ALIGN_PARAGRAPH.CENTER
    university.paragraph_format.space_after = Pt(2)
    set_font(university.add_run("TRƯỜNG ĐẠI HỌC BÌNH DƯƠNG"), size=13, bold=True, color=NAVY)
    institute = doc.add_paragraph()
    institute.alignment = WD_ALIGN_PARAGRAPH.CENTER
    institute.paragraph_format.space_after = Pt(2)
    set_font(institute.add_run("VIỆN TRÍ TUỆ NHÂN TẠO VÀ CHUYỂN ĐỔI SỐ"), size=11, bold=True, color=NAVY)
    faculty = doc.add_paragraph()
    faculty.alignment = WD_ALIGN_PARAGRAPH.CENTER
    faculty.paragraph_format.space_after = Pt(70)
    set_font(faculty.add_run("KHOA CÔNG NGHỆ THÔNG TIN, ROBOT VÀ TRÍ TUỆ NHÂN TẠO"), size=10.5, bold=True, color=NAVY)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_font(
        kicker.add_run("BÁO CÁO ĐỒ ÁN TỐT NGHIỆP"),
        size=11,
        bold=True,
        color=BLUE,
    )

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("BOOKVERSE AI"), size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    set_font(
        subtitle.add_run("Nền tảng đọc và giao dịch sách tích hợp"),
        size=14,
        bold=True,
        color=DARK_BLUE,
    )
    subtitle2 = doc.add_paragraph()
    subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle2.paragraph_format.space_after = Pt(28)
    set_font(
        subtitle2.add_run("hệ gợi ý có giải thích và trợ lý RAG kiểm soát nguồn"),
        size=14,
        bold=True,
        color=DARK_BLUE,
    )

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(45)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    p_pr.append(borders)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta.paragraph_format.left_indent = Inches(1.45)
    meta.paragraph_format.space_after = Pt(8)
    set_font(meta.add_run("Sinh viên thực hiện:  LƯƠNG NGUYỄN QUỐC TUẤN"), size=12, bold=True, color=NAVY)
    mssv = doc.add_paragraph()
    mssv.alignment = WD_ALIGN_PARAGRAPH.LEFT
    mssv.paragraph_format.left_indent = Inches(1.45)
    set_font(mssv.add_run("Mã số sinh viên:       22050098"), size=12, bold=True, color=NAVY)

    year = doc.add_paragraph()
    year.paragraph_format.space_before = Pt(78)
    year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(year.add_run("Thành phố Hồ Chí Minh, tháng 8 năm 2026"), size=11, italic=True, color=NAVY)
    doc.add_page_break()


def add_acknowledgements(doc: Document) -> None:
    heading = doc.add_paragraph("LỜI CẢM ƠN", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs = [
        "Trong quá trình học tập và thực hiện đồ án tốt nghiệp, em đã nhận được sự quan tâm, hướng dẫn và động viên từ quý thầy cô, gia đình và bạn bè.",
        "Em xin trân trọng cảm ơn Ban Giám hiệu Trường Đại học Bình Dương, Viện Trí tuệ nhân tạo và Chuyển đổi số cùng quý thầy cô Khoa Công nghệ Thông tin, Robot và Trí tuệ nhân tạo đã trang bị cho em kiến thức nền tảng và tạo điều kiện để em hoàn thành đồ án.",
        "Em cũng xin cảm ơn các bạn đã góp ý trong quá trình kiểm thử và hoàn thiện sản phẩm. Dù đã cố gắng, báo cáo vẫn có thể còn những hạn chế; em mong nhận được các ý kiến đóng góp để tiếp tục cải thiện trong thời gian tới.",
        "Em xin chân thành cảm ơn!",
    ]
    for text in paragraphs:
        paragraph_from_text(doc, text)
    signature = doc.add_paragraph()
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signature.paragraph_format.space_before = Pt(20)
    set_font(signature.add_run("Sinh viên thực hiện\nLương Nguyễn Quốc Tuấn"), size=11, italic=True, color=NAVY)
    doc.add_page_break()


def _font(size: int, bold: bool = False):
    """Chọn font có sẵn trên Windows để chữ tiếng Việt trong sơ đồ không bị lỗi."""
    path = "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"
    return ImageFont.truetype(path, size)


def _centered(draw, box, text: str, font, fill: str) -> None:
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=6)
    x = (left + right - (bbox[2] - bbox[0])) / 2
    y = (top + bottom - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), text, font=font, fill=fill, align="center", spacing=6)


def _box(draw, box, text: str, fill: str, border: str = "#2E74B5") -> None:
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=border, width=3)
    _centered(draw, box, text, _font(27, True), "#0B2545")


def _arrow(draw, start, end) -> None:
    draw.line([start, end], fill="#526777", width=5)
    draw.polygon([(end[0], end[1]), (end[0] - 16, end[1] - 10), (end[0] - 16, end[1] + 10)], fill="#526777")


def build_diagrams() -> dict[str, Path]:
    """Tạo hai sơ đồ gốc từ kiến trúc thật của BookVerse để chèn vào báo cáo."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    image = Image.new("RGB", (1800, 920), "white")
    draw = ImageDraw.Draw(image)
    _centered(draw, (0, 25, 1800, 95), "KIẾN TRÚC TỔNG QUAN BOOKVERSE AI", _font(36, True), "#0B2545")
    boxes = [
        ((70, 225, 390, 415), "Người dùng\nĐộc giả | Người bán | Quản trị", "#F4F6F9"),
        ((540, 225, 875, 415), "Next.js 15\nGiao diện và nghiệp vụ", "#E8EEF5"),
        ((1030, 125, 1370, 315), "PostgreSQL + Prisma\nDữ liệu và giao dịch", "#E8EEF5"),
        ((1030, 505, 1370, 695), "FastAPI\nDịch vụ gợi ý AI", "#E8EEF5"),
        ((1490, 315, 1740, 505), "LLM tùy chọn\nTrợ lý RAG", "#F4F6F9"),
    ]
    for box, text, color in boxes:
        _box(draw, box, text, color)
    _arrow(draw, (390, 320), (540, 320)); _arrow(draw, (875, 270), (1030, 220)); _arrow(draw, (875, 370), (1030, 600)); _arrow(draw, (1370, 410), (1490, 410))
    path = ASSET_DIR / "so_do_kien_truc_bookverse.png"
    image.save(path); paths["architecture"] = path

    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    _centered(draw, (0, 25, 1800, 95), "LUỒNG KIỂM SOÁT QUYỀN ĐỌC SÁCH", _font(36, True), "#0B2545")
    boxes = [
        ((70, 275, 330, 465), "Người dùng\nmở nội dung", "#F4F6F9"),
        ((470, 275, 760, 465), "Xác thực\nphiên đăng nhập", "#E8EEF5"),
        ((900, 275, 1200, 465), "Kiểm tra\nentitlement", "#E8EEF5"),
        ((1340, 130, 1710, 320), "Có quyền\nTrả nội dung hợp lệ", "#DDEBF7"),
        ((1340, 470, 1710, 660), "Chưa có quyền\nChỉ trả phần xem thử", "#FCE4D6"),
    ]
    for box, text, color in boxes:
        _box(draw, box, text, color)
    _arrow(draw, (330, 370), (470, 370)); _arrow(draw, (760, 370), (900, 370)); _arrow(draw, (1200, 325), (1340, 225)); _arrow(draw, (1200, 415), (1340, 565))
    path = ASSET_DIR / "so_do_quyen_doc.png"
    image.save(path); paths["entitlement"] = path
    return paths


def add_figure(doc: Document, path: Path, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    picture = paragraph.add_run().add_picture(str(path), width=Inches(6.35))
    # Mô tả thay thế giúp trình đọc màn hình nhận diện đúng nội dung sơ đồ.
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", caption)
    doc_pr.set("title", caption)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(10)
    set_font(caption_p.add_run(caption), size=10, italic=True, color=MUTED)


def add_toc(doc: Document) -> None:
    heading = doc.add_paragraph("MỤC LỤC", style="Heading 1")
    heading.paragraph_format.page_break_before = False
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(
        note.add_run(
            "Mục lục được cập nhật tự động khi mở tài liệu trong Microsoft Word."
        ),
        size=9.5,
        italic=True,
        color=MUTED,
    )
    toc = doc.add_paragraph()
    add_field(toc, r'TOC \o "1-3" \h \z \u', "Nhấn Ctrl+A, F9 để cập nhật mục lục.")
    doc.add_page_break()


def paragraph_from_text(doc: Document, text: str, *, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.widow_control = True
    add_inline(paragraph, text)
    return paragraph


def create_numbering_instance(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "280")
    p_pr.extend([tabs, indent])
    level.extend([start, fmt, text, suffix, p_pr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_numbered_items(doc: Document, items: list[str]) -> None:
    num_id = create_numbering_instance(doc)
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.208
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, number])
        p_pr.append(num_pr)
        add_inline(paragraph, item)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GRAY)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    add_inline(paragraph, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table_widths(rows: list[list[str]]) -> list[int]:
    columns = max(len(row) for row in rows)
    lengths = []
    for index in range(columns):
        max_length = max(len(row[index]) if index < len(row) else 0 for row in rows)
        lengths.append(max(8, min(max_length, 42)))
    total = sum(lengths)
    widths = [round(9360 * length / total) for length in lengths]
    widths[-1] += 9360 - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, source_row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            text = source_row[column_index] if column_index < len(source_row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(8.5 if columns >= 5 else 9.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = NAVY
            if row_index == 0:
                shade_cell(cell, LIGHT_BLUE)
    set_table_geometry(table, table_widths(rows))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_markdown_body(doc: Document, source: str, diagrams: dict[str, Path]) -> None:
    lines = source.splitlines()
    start = next(index for index, line in enumerate(lines) if line == "## Tóm tắt")
    index = start
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph_from_text(doc, " ".join(line.strip() for line in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            flush_paragraph()
            index += 1
            continue

        if line == "<!-- PAGE_BREAK -->":
            flush_paragraph()
            doc.add_page_break()
            index += 1
            continue

        if line.startswith("|"):
            flush_paragraph()
            block: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip())
                index += 1
            parsed = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in block
                if not re.match(r"^\|[\s:|-]+\|$", row)
            ]
            add_markdown_table(doc, parsed)
            continue

        if line.startswith("#"):
            flush_paragraph()
            level = len(line) - len(line.lstrip("#"))
            title = line[level:].strip()
            if level == 1:
                paragraph = doc.add_paragraph(title, style="Heading 1")
                if title.startswith("Chương "):
                    paragraph.paragraph_format.page_break_before = True
            elif level == 2:
                paragraph = doc.add_paragraph(title, style="Heading 2")
            else:
                paragraph = doc.add_paragraph(title, style="Heading 3")
            index += 1
            if title.startswith("4.1 Kiến trúc"):
                add_figure(doc, diagrams["architecture"], "Hình 4.1. Sơ đồ kiến trúc tổng quan của hệ thống BookVerse AI")
            elif title.startswith("4.3 Luồng quyền đọc"):
                add_figure(doc, diagrams["entitlement"], "Hình 4.2. Sơ đồ kiểm soát quyền truy cập nội dung đọc")
            continue

        if line.startswith("> "):
            flush_paragraph()
            add_callout(doc, line[2:].strip())
            index += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            flush_paragraph()
            items: list[str] = []
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index]):
                items.append(re.sub(r"^\d+\.\s+", "", lines[index]).strip())
                index += 1
            add_numbered_items(doc, items)
            continue

        if line.startswith("- "):
            flush_paragraph()
            paragraph_from_text(doc, line[2:].strip(), style="List Bullet")
            index += 1
            continue

        paragraph_buffer.append(line)
        index += 1
    flush_paragraph()


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    source = SOURCE.read_text(encoding="utf-8")
    diagrams = build_diagrams()
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    add_running_furniture(doc.sections[0])
    add_cover(doc)
    add_acknowledgements(doc)
    add_toc(doc)
    parse_markdown_body(doc, source, diagrams)

    # Bật cập nhật field khi mở bằng Word.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.core_properties.title = "BookVerse AI - Báo cáo đồ án tốt nghiệp"
    doc.core_properties.subject = "Hệ gợi ý có giải thích và trợ lý RAG kiểm soát nguồn"
    doc.core_properties.author = "Lương Nguyễn Quốc Tuấn - 22050098"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
