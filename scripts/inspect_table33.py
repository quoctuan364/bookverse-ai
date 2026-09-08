import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_HOAN_THIEN_CHUONG_5_KET_LUAN.docx')

t = doc.tables[33]
for row_idx, r in enumerate(t.rows):
    print(f'=== Row {row_idx} ===')
    for cell_idx, c in enumerate(r.cells):
        print(f'  Cell {cell_idx}:')
        for p in c.paragraphs:
            print(f'    P (style: {p.style.name}): {p.text}')
            for run in p.runs:
                print(f'      Run: "{run.text}" | bold={run.bold}, font={run.font.name}, size={run.font.size}')
