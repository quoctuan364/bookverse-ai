# -*- coding: utf-8 -*-
"""Cập nhật luận văn theo góp ý về quy trình dữ liệu và benchmark.

Script luôn đọc bản DOCX hiện tại và ghi sang một tệp mới, không ghi đè bản gốc.
Các số liệu được lấy từ artifact kiểm thử đã có trong repository.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table


ROOT = Path(r"D:\Doantotnghiep")
SOURCE_DOCX = ROOT / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098.docx"
OUTPUT_DOCX = ROOT / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_CAP_NHAT_GOP_Y_DATASET.docx"
ASSET_DIR = ROOT / "outputs" / "thesis_revision_assets"
PIPELINE_IMAGE = ASSET_DIR / "dataset_pipeline.png"
EVALUATION_IMAGE = ASSET_DIR / "evaluation_actual.png"

FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def draw_centered_multiline(draw, box, text, font, fill, spacing=5):
    """Vẽ văn bản nhiều dòng vào giữa một hộp."""
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(
        ((x1 + x2 - width) / 2, (y1 + y2 - height) / 2),
        text,
        font=font,
        fill=fill,
        spacing=spacing,
        align="center",
    )


def build_pipeline_image() -> None:
    """Tạo sơ đồ pipeline dữ liệu bằng ảnh raster rõ nét để nhúng vào Word."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 780), "white")
    draw = ImageDraw.Draw(image)
    title_font = pil_font(42, bold=True)
    box_title_font = pil_font(28, bold=True)
    box_body_font = pil_font(24)
    draw.text((900, 55), "QUY TRÌNH THU THẬP VÀ KIỂM SOÁT DATASET BOOKVERSE AI",
              font=title_font, fill="#173F5F", anchor="mm")

    boxes = [
        (60, 180, 320, 600, "1. NGUỒN", "Google Books API\nOpen Library API\nDataset synthetic"),
        (400, 180, 660, 600, "2. THU THẬP", "12 truy vấn tiếng Việt\nSearch/Works API\nRate limit + retry"),
        (740, 180, 1000, 600, "3. CHUẨN HÓA", "Title, author, ISBN\nWork/edition key\nCategory canonical"),
        (1080, 180, 1340, 600, "4. QUALITY GATE", "Loại trùng\nKiểm tra provenance\nCover policy + checksum"),
        (1420, 180, 1680, 600, "5. SỬ DỤNG", "Catalog demo\nTemporal split\nBenchmark read-only"),
    ]
    fills = ["#E8F3F1", "#EAF0F8", "#FFF4D6", "#FCE9E4", "#EDE8F7"]
    for index, (x1, y1, x2, y2, heading, body) in enumerate(boxes):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=fills[index], outline="#315F64", width=4)
        draw_centered_multiline(draw, (x1 + 15, y1 + 35, x2 - 15, y1 + 145), heading,
                                box_title_font, "#173F5F")
        draw.line((x1 + 35, y1 + 155, x2 - 35, y1 + 155), fill="#8BA9A7", width=3)
        draw_centered_multiline(draw, (x1 + 20, y1 + 175, x2 - 20, y2 - 25), body,
                                box_body_font, "#263238", spacing=13)
        if index < len(boxes) - 1:
            start_x = x2 + 18
            end_x = boxes[index + 1][0] - 18
            middle_y = (y1 + y2) // 2
            draw.line((start_x, middle_y, end_x, middle_y), fill="#D47B5B", width=7)
            draw.polygon([(end_x, middle_y), (end_x - 23, middle_y - 16), (end_x - 23, middle_y + 16)],
                         fill="#D47B5B")
    draw.text((900, 700),
              "Nguyên tắc: không ghi đè dữ liệu gốc - tách rõ metadata thật và tương tác synthetic - mọi benchmark có checksum",
              font=pil_font(25), fill="#4E5D63", anchor="mm")
    image.save(PIPELINE_IMAGE, quality=95)


def build_evaluation_image() -> None:
    """Tạo biểu đồ từ metric temporal evaluation thực tế tại K=10."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    methods = ["Popularity", "Content", "Behavior", "Hybrid\nproduction", "Random\nsanity"]
    ndcg = [0.003516, 0.006419, 0.005374, 0.001496, 0.002367]
    hit_rate = [0.013655, 0.019958, 0.021008, 0.008403, 0.010504]
    image = Image.new("RGB", (1700, 930), "white")
    draw = ImageDraw.Draw(image)
    draw.text((850, 55), "KẾT QUẢ TEMPORAL EVALUATION THỰC TẾ TẠI K = 10",
              font=pil_font(42, bold=True), fill="#173F5F", anchor="mm")
    plot = (145, 150, 1600, 760)
    x1, y1, x2, y2 = plot
    max_value = 0.024
    for tick in range(0, 25, 4):
        value = tick / 1000
        y = y2 - (value / max_value) * (y2 - y1)
        draw.line((x1, y, x2, y), fill="#D7DEE2", width=2)
        draw.text((x1 - 20, y), f"{value:.3f}", font=pil_font(21), fill="#56636A", anchor="rm")
    draw.line((x1, y1, x1, y2), fill="#34495E", width=4)
    draw.line((x1, y2, x2, y2), fill="#34495E", width=4)

    group_width = (x2 - x1) / len(methods)
    bar_width = 70
    for i, method in enumerate(methods):
        center = x1 + group_width * (i + 0.5)
        for offset, value, color in [(-42, ndcg[i], "#2A7F78"), (42, hit_rate[i], "#D77A61")]:
            left = center + offset - bar_width / 2
            right = center + offset + bar_width / 2
            top = y2 - (value / max_value) * (y2 - y1)
            draw.rounded_rectangle((left, top, right, y2), radius=8, fill=color)
            draw.text(((left + right) / 2, top - 10), f"{value:.4f}", font=pil_font(19), fill="#263238", anchor="ms")
        draw.multiline_text((center, y2 + 35), method, font=pil_font(22, bold=True),
                            fill="#263238", anchor="ma", align="center", spacing=4)
    draw.rounded_rectangle((550, 845, 600, 885), radius=6, fill="#2A7F78")
    draw.text((620, 865), "NDCG@10", font=pil_font(24), fill="#263238", anchor="lm")
    draw.rounded_rectangle((880, 845, 930, 885), radius=6, fill="#D77A61")
    draw.text((950, 865), "HitRate@10", font=pil_font(24), fill="#263238", anchor="lm")
    image.save(EVALUATION_IMAGE, quality=95)


def set_run_font(run, size=13, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def style_paragraph(paragraph, *, size=13, bold=False, italic=False,
                    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True,
                    before=0, after=4, line_spacing=1.3):
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.first_line_indent = Cm(1.27) if indent else None
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def replace_paragraph(paragraph, text, **style):
    paragraph.text = text
    return style_paragraph(paragraph, **style)


def find_paragraph(doc, exact_text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"Không tìm thấy đoạn: {exact_text}")


def find_paragraph_contains(doc, text_fragment):
    for paragraph in doc.paragraphs:
        if text_fragment in paragraph.text:
            return paragraph
    raise ValueError(f"Không tìm thấy đoạn chứa: {text_fragment}")


def table_before_paragraph(paragraph):
    """Lấy bảng gần nhất nằm trước một chú thích, tránh phụ thuộc chỉ số bảng."""
    element = paragraph._p.getprevious()
    while element is not None:
        if element.tag == qn("w:tbl"):
            return Table(element, paragraph._parent)
        element = element.getprevious()
    raise ValueError(f"Không tìm thấy bảng trước chú thích: {paragraph.text}")


def insert_paragraph(anchor, text, *, kind="body"):
    style_name = None
    if kind == "number":
        style_name = "List Number"
    elif kind == "bullet":
        style_name = "List Bullet"
    paragraph = anchor.insert_paragraph_before(text, style=style_name)
    if kind == "h1":
        return style_paragraph(paragraph, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                               indent=False, before=10, after=4)
    if kind == "h2":
        return style_paragraph(paragraph, bold=True, italic=True,
                               alignment=WD_ALIGN_PARAGRAPH.LEFT, indent=False,
                               before=7, after=3)
    if kind in {"number", "bullet"}:
        paragraph.paragraph_format.left_indent = Cm(1.27)
        paragraph.paragraph_format.first_line_indent = None
        return style_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                               indent=False, after=3)
    if kind == "caption":
        return style_paragraph(paragraph, size=11.5, bold=True, italic=True,
                               alignment=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                               before=4, after=8)
    return style_paragraph(paragraph)


def set_cell_margins(cell, top=90, bottom=90, left=110, right=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, header=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5):
    cell.text = str(text)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    shade_cell(cell, "1F4E79" if header else "FFFFFF")
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=header, color=(255, 255, 255) if header else None)


def fill_table(table, headers, rows, widths=None):
    """Thay toàn bộ nội dung bảng nhưng giữ đối tượng bảng tại vị trí cũ."""
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.columns) < len(headers):
        table.add_column(Cm(2))
    while len(table.columns) > len(headers):
        for row in table.rows:
            row._tr.remove(row.cells[-1]._tc)
    header_row = table.rows[0]
    for index, header in enumerate(headers):
        set_cell_text(header_row.cells[index], header, header=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    tr_pr = header_row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            numeric = column_index > 0 and len(str(value)) < 18
            set_cell_text(cells[column_index], value, align=WD_ALIGN_PARAGRAPH.CENTER if numeric else WD_ALIGN_PARAGRAPH.LEFT)
            shade_cell(cells[column_index], "F4F7FB" if row_index % 2 else "FFFFFF")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)


def insert_table(doc, anchor, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    fill_table(table, headers, rows, widths)
    anchor._p.addprevious(table._tbl)
    return table


def insert_dataset_section(doc) -> None:
    anchor = find_paragraph(doc, "2.3. Công nghệ sử dụng trong BookVerse AI :")
    insert_paragraph(anchor, "2.2. Quy trình thu thập và xây dựng dataset", kind="h1")
    insert_paragraph(
        anchor,
        "Dữ liệu của BookVerse AI được tổ chức thành hai lớp tách biệt để tránh đánh đồng metadata thư mục với hành vi người dùng. Lớp synthetic dùng để kiểm tra pipeline và temporal evaluation; lớp bibliographic metadata được tuyển chọn từ API công khai để trình diễn catalog và kiểm tra provenance. Không lớp nào được tuyên bố là dữ liệu hành vi production.",
    )
    insert_table(
        doc,
        anchor,
        ["Lớp dữ liệu", "Nguồn", "Quy mô đã xác minh", "Nhãn provenance", "Mục đích"],
        [
            ["Ultra Synthetic", "Script sinh dữ liệu nội bộ", "2.200 sách; 18.000 sự kiện", "SYNTHETIC_DATA", "Kiểm tra pipeline và benchmark temporal"],
            ["Real Curated", "Open Library API", "3.046 record = 3.044 work + 2 edition-only", "CURATED_REAL_BIBLIOGRAPHIC_METADATA", "Catalog, kiểm tra nguồn và giao diện"],
            ["Google Books demo", "Google Books API", "Thu theo batch; chỉ giữ bản ghi hợp lệ", "DEMO_METADATA", "Bổ sung dữ liệu trình diễn, không dùng benchmark"],
        ],
        [3.0, 3.0, 4.0, 3.8, 4.2],
    )
    insert_paragraph(anchor, "Bảng 2-1: Kiến trúc dữ liệu hai lớp và phạm vi sử dụng", kind="caption")

    insert_paragraph(anchor, "2.2.1. Kênh Google Books API", kind="h2")
    google_steps = [
        "Xác định 12 truy vấn tiếng Việt gồm tiểu thuyết Việt Nam, kinh doanh, tâm lý học, lịch sử Việt Nam, kỹ năng sống, khoa học, giáo dục, văn học Việt Nam, thiếu nhi, khởi nghiệp, quản trị và công nghệ.",
        "Gọi GET /books/v1/volumes với q, langRestrict=vi, printType=books, maxResults=40 và startIndex; mỗi từ khóa lấy tối đa 3 trang, timeout 20 giây và nghỉ 0,3 giây giữa các request [10].",
        "Loại bản ghi thiếu title hoặc cover; lấy tác giả đầu tiên; loại HTML và chuẩn hóa khoảng trắng trong description; chuyển URL ảnh từ HTTP sang HTTPS.",
        "Khử trùng theo khóa title.lower() + author.lower(). Giá và category sinh cho demo được gắn nhãn SYNTHETIC_DEMO_PRICE và DEMO_CATEGORY, không coi là metadata thị trường.",
        "Ghi UTF-8 vào data/demo/books.csv. Script tạo books.csv.bak một lần trước khi ghi, vì vậy bản dữ liệu gốc có đường phục hồi.",
    ]
    for step in google_steps:
        insert_paragraph(anchor, step, kind="number")

    insert_paragraph(anchor, "2.2.2. Kênh Open Library API", kind="h2")
    open_steps = [
        "Thu metadata theo nhóm chủ đề bằng Search API và đối chiếu Work/Edition API. Các trường gồm source key, title, author, subject, năm xuất bản đầu tiên, ISBN, cover ID, mô tả và language [11].",
        "Chuẩn hóa source identity về /works/OL...W; nếu chỉ xác minh được edition thì giữ nhãn EDITION_ONLY thay vì tự suy diễn work key.",
        "Kiểm tra định dạng và checksum ISBN; loại trùng source key/ISBN; ánh xạ 39 category nguồn vào 27 nhóm canonical; giá vẫn được ghi rõ là SYNTHETIC_DEMO_PRICE.",
        "Audit cover theo allowlist covers.openlibrary.org, kiểm tra redirect, content type, kích thước và tỷ lệ. HTTP hợp lệ không đồng nghĩa có quyền tái phân phối; quyền ảnh giữ trạng thái NOT_VERIFIED.",
        "Pipeline chạy tuần tự validate-only, dry-run và execute. Lượt execute thứ hai phải idempotent: 0 insert, 0 update và 3.046 unchanged trên database test.",
    ]
    for step in open_steps:
        insert_paragraph(anchor, step, kind="number")
    insert_table(
        doc,
        anchor,
        ["Chỉ số catalog Open Library", "Giá trị đã xác minh"],
        [
            ["Tổng source record", "3.046"],
            ["Work / Edition-only", "3.044 / 2"],
            ["ISBN duy nhất", "2.343"],
            ["Tác giả duy nhất", "3.260"],
            ["Ấn bản tiếng Việt", "259"],
            ["Thiếu language", "34 - giữ NOT_AVAILABLE"],
            ["Bản ghi OCR/chuyển mã bị loại", "414"],
            ["Cover đạt technical policy", "687/3.046; quyền ảnh NOT_VERIFIED"],
        ],
        [9.0, 9.0],
    )
    insert_paragraph(anchor, "Bảng 2-2: Thống kê chất lượng catalog bibliographic tuyển chọn", kind="caption")

    insert_paragraph(anchor, "2.2.3. Chuẩn hóa, kiểm soát chất lượng và provenance", kind="h2")
    insert_paragraph(
        anchor,
        "Quality gate được triển khai theo nguyên tắc fail-closed: bản ghi thiếu trường bắt buộc hoặc không chứng minh được source identity sẽ bị reject/đưa vào hàng chờ review; dữ liệu nguồn không bị sửa tại chỗ. Mỗi artifact sinh ra có report, checksum SHA-256 và trạng thái VERIFIED, PARTIAL hoặc NOT_AVAILABLE. Open Library khuyến nghị API cho truy vấn thời gian thực, lưu lượng thấp và sử dụng cache; vì vậy pipeline giới hạn tốc độ, retry có backoff và không dùng API làm backend tải hàng loạt [11].",
    )
    picture_paragraph = anchor.insert_paragraph_before()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(PIPELINE_IMAGE), width=Cm(17.2))
    insert_paragraph(anchor, "Hình 2-1: Pipeline thu thập, chuẩn hóa và kiểm soát dataset BookVerse AI", kind="caption")

    insert_paragraph(anchor, "2.2.4. Tạo ground truth và chia tập thực nghiệm", kind="h2")
    insert_paragraph(
        anchor,
        "Benchmark chỉ dùng lớp synthetic 2.200 sách; catalog Open Library không được trộn vào temporal evaluation cũ. Strong-positive gồm purchase có đơn hợp lệ và quantity dương, ReadingSession từ 300 giây hoặc progressPercent từ 50%, Bookmark, Favorite và Review từ 4 sao. VIEW, SEARCH, CART, COMMENT, REACTION, HIGHLIGHT, CHATBOT_QUERY, order PENDING/CANCELLED/REFUNDED và PURCHASE event thiếu orderId đều bị loại.",
    )
    insert_paragraph(
        anchor,
        "Mốc global temporal cutoff là 01/06/2026: train có 12.206 positive trước cutoff, test có 2.593 positive từ cutoff trở đi. Candidate snapshot gồm 2.000 sách ACTIVE tạo trước cutoff; 952 user đủ điều kiện và 2.349 cặp user-book ground truth. Cách chia theo thời gian giúp hạn chế việc mô hình học từ tương tác tương lai - một nguồn data leakage phổ biến trong offline recommendation evaluation [12].",
    )
    insert_table(
        doc,
        anchor,
        ["Thành phần benchmark", "Số lượng"],
        [
            ["Book nguồn synthetic", "2.200"],
            ["InteractionEvent", "18.000"],
            ["ReadingSession / Bookmark / Favorite", "6.200 / 2.799 / 48"],
            ["Review / OrderItem", "3.600 / 4.714"],
            ["Positive sau policy", "14.799"],
            ["Train / Test positive", "12.206 / 2.593"],
            ["User đủ điều kiện", "952"],
        ],
        [10.0, 8.0],
    )
    insert_paragraph(anchor, "Bảng 2-3: Thành phần dataset dùng cho temporal evaluation", kind="caption")

    insert_paragraph(anchor, "2.2.5. Giới hạn và nguyên tắc diễn giải", kind="h2")
    limits = [
        "18.000 sự kiện là synthetic, không đại diện hành vi người dùng thật; CTR và UAT/SUS hiện là NOT_AVAILABLE.",
        "3.046 record Open Library là metadata thư mục, không kèm interaction, order hoặc review thật; 2 record chỉ có edition key.",
        "Giá demo và category demo không được dùng để suy ra giá thị trường hoặc chất lượng mô hình.",
        "Metric offline chỉ phản ánh pipeline trên snapshot đã khóa; không được diễn giải thành hiệu quả production hay tăng doanh thu.",
        "Khi thu được dữ liệu người dùng thật có consent, cần tạo split mới và benchmark lại trước khi thay đổi trọng số production.",
    ]
    for item in limits:
        insert_paragraph(anchor, item, kind="bullet")


def revise_existing_content(doc) -> None:
    replace_paragraph(find_paragraph(doc, "Tính mới và tính mới và giải pháp của BookVerse AI:"),
                      "Tính mới và giải pháp của BookVerse AI:", indent=False)
    replace_paragraph(
        find_paragraph(doc, "• Xây dựng website BookVerse AI hoàn chỉnh trên nền tảng Next.js 15, Prisma ORM, PostgreSQL pgvector và Python FastAPI."),
        "• Xây dựng website BookVerse AI trên nền tảng Next.js 15, Prisma ORM, PostgreSQL/pgvector và Python FastAPI; kiểm chứng bằng test tự động và build production.",
        indent=False,
    )
    replace_paragraph(
        find_paragraph(doc, "• Triển khai hệ thống gợi ý lai (Hybrid Recommender System) đạt độ chính xác cao (F1-score > 87%) và thời gian phản hồi < 1.5s."),
        "• Triển khai và đánh giá hệ thống gợi ý bằng temporal split, so sánh Popularity, Content, Behavior và Hybrid theo Precision, Recall, HitRate, NDCG, MRR và Coverage; không đặt trước kết quả phải đạt.",
        indent=False,
    )
    # Đề cương đầu tài liệu nằm trong bảng nên không xuất hiện trong doc.paragraphs.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "Hệ thống đề xuất sách thông minh đạt độ chính xác cao" in paragraph.text:
                        replace_paragraph(
                            paragraph,
                            "• Hệ thống gợi ý được đánh giá bằng temporal split và các chỉ số Precision@K, Recall@K, HitRate@K, NDCG@K; kết quả được báo cáo theo artifact thực nghiệm, không đặt trước mức phải đạt.",
                            indent=False,
                        )
    replace_paragraph(
        find_paragraph(doc, "Trong BookVerse AI, mỗi cuốn sách được chuyển đổi thành một vector đặc trưng nhiều chiều (Dense Vector Embedding) biểu diễn toàn diện ngữ nghĩa của tên sách, tác giả, thể loại và tóm tắt nội dung. Việc tính toán độ tương đồng giữa hai vector đặc trưng được thực hiện trực tiếp trong cơ sở dữ liệu PostgreSQL nhờ extension pgvector thông qua khoảng cách Cosine:"),
        "Schema BookVerse AI hỗ trợ lưu vector embedding và truy vấn khoảng cách cosine bằng toán tử <=> của pgvector. Bản triển khai hiện dùng kiểu vector chưa khóa số chiều và chưa tạo HNSW/IVFFlat; vì vậy tài liệu chỉ mô tả khả năng tìm kiếm exact hiện có, còn ANN index được xếp vào hướng tối ưu khi quy mô tăng [7].",
    )

    technology_caption = find_paragraph(doc, "Bảng 2-1: Bảng công nghệ sử dụng trong BookVerse AI")
    for technology_table in doc.tables:
        for row in technology_table.rows:
            first_cell = row.cells[0].text.strip()
            if "pgvector" in first_cell and ("PostgreSQL" in first_cell or first_cell == "pgvector"):
                set_cell_text(row.cells[0], "PostgreSQL pgvector")
                set_cell_text(row.cells[1], "Lưu kiểu vector và hỗ trợ toán tử khoảng cách cosine; hiện chưa tạo chỉ mục ANN.")
                set_cell_text(row.cells[2], "Phục vụ truy vấn vector exact; HNSW/IVFFlat là hướng tối ưu có điều kiện khi dữ liệu tăng.")
    replace_paragraph(
        technology_caption,
        "Bảng 2-4: Bảng công nghệ sử dụng trong BookVerse AI",
        size=11.5, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False, before=4, after=8,
    )

    # Mục lục của bản gốc là văn bản tĩnh, vì vậy cập nhật thủ công các dòng Chương 2.
    toc_replacements = [
        ("2.2. Kiến trúc Web Fullstack hiện đại", "  2.2. Quy trình thu thập và xây dựng dataset ........................ 9"),
        ("2.2.1. Mô hình App Router", "    2.2.1. Kênh Google Books API ..................................... 10"),
        ("2.2.2. Cơ chế lưu trữ Vector", "    2.2.2. Kênh Open Library API ..................................... 10"),
        ("2.2.3. Quy trình phát triển hệ thống", "    2.2.3. Chuẩn hóa, kiểm soát chất lượng và provenance ............. 11"),
    ]
    for text_fragment, new_text in toc_replacements:
        replace_paragraph(find_paragraph_contains(doc, text_fragment), new_text, size=12,
                          alignment=WD_ALIGN_PARAGRAPH.LEFT, indent=False,
                          before=0, after=0, line_spacing=1.0)
    toc_anchor = find_paragraph_contains(doc, "2.3. Công nghệ sử dụng trong BookVerse AI")
    for text in [
        "    2.2.4. Tạo ground truth và chia tập thực nghiệm ...................... 12",
        "    2.2.5. Giới hạn và nguyên tắc diễn giải .............................. 13",
    ]:
        paragraph = toc_anchor.insert_paragraph_before(text)
        style_paragraph(paragraph, size=12, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                        indent=False, before=0, after=0, line_spacing=1.0)

    evaluation_rows = [
        ["Popularity", "10", "0,001366", "0,006197", "0,013655", "0,003516"],
        ["Content-based", "10", "0,002416", "0,010812", "0,019958", "0,006419"],
        ["Behavior-based", "10", "0,002521", "0,010530", "0,021008", "0,005374"],
        ["Hybrid production", "10", "0,000840", "0,002451", "0,008403", "0,001496"],
        ["Random seeded sanity", "10", "0,001050", "0,005252", "0,010504", "0,002367"],
    ]
    evaluation_caption = find_paragraph(doc, "Bảng 4-1: Kết quả thử nghiệm hiệu năng các mô hình trong BookVerse AI")
    fill_table(
        table_before_paragraph(evaluation_caption),
        ["Phương pháp", "K", "Precision", "Recall", "Hit Rate", "NDCG"],
        evaluation_rows,
        [4.8, 1.3, 2.8, 2.8, 3.0, 2.8],
    )
    replace_paragraph(evaluation_caption,
                      "Bảng 4-1: Kết quả temporal evaluation trên cùng cohort tại K = 10",
                      size=11.5, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                      indent=False, before=4, after=8)

    # Thay ảnh biểu đồ cũ bằng biểu đồ metric có nguồn từ artifact evaluation.
    figure_caption = find_paragraph(doc, "Hình 4-1: Biểu đồ so sánh hiệu năng các mô hình thực nghiệm trong BookVerse AI")
    previous = figure_caption._p.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        for child in list(previous):
            if child.tag != qn("w:pPr"):
                previous.remove(child)
        from docx.text.paragraph import Paragraph
        picture_paragraph = Paragraph(previous, figure_caption._parent)
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.add_run().add_picture(str(EVALUATION_IMAGE), width=Cm(16.8))
    replace_paragraph(figure_caption,
                      "Hình 4-1: NDCG@10 và HitRate@10 từ temporal evaluation; Hybrid chưa vượt baseline",
                      size=11.5, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                      indent=False, before=4, after=8)

    exception_caption = find_paragraph(doc, "Bảng 4-4: Bảng xử lý các trường hợp ngoại lệ trong BookVerse AI")
    fill_table(
        table_before_paragraph(exception_caption),
        ["Trường hợp kiểm tra", "Bằng chứng", "Cách xử lý", "Kết luận", "Giới hạn diễn giải"],
        [
            ["User cold-start", "12 user không có train feature", "Fallback deterministic; báo cohort riêng", "Có đường phục hồi", "Không đủ để tuyên bố cá nhân hóa"],
            ["User sparse", "1 user có 1-2 sách", "Tách khỏi warm cohort", "Không kết luận rộng", "Cỡ mẫu quá nhỏ"],
            ["Positive ngoài candidate", "220 cặp", "Loại khỏi ground truth", "Tránh chấm mô hình bằng item chưa hợp lệ", "Chỉ đúng snapshot cutoff"],
            ["Item đã thấy trong train", "19 cặp", "Loại khỏi ground truth", "Tránh leakage theo user", "Không đo novelty online"],
            ["Evidence thiếu provenance", "8.400/8.400", "UI dùng nhãn trung tính", "Không tạo claim REAL_USER_DATA", "CTR/UAT vẫn NOT_AVAILABLE"],
        ],
        [3.2, 3.2, 4.4, 3.6, 4.6],
    )
    replace_paragraph(exception_caption,
                      "Bảng 4-2: Kiểm soát các trường hợp biên trong pipeline đánh giá",
                      size=11.5, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                      indent=False, before=4, after=8)

    goal_caption = find_paragraph(doc, "Bảng 5-1: Kết quả đối chiếu với mục tiêu đề tài BookVerse AI")
    fill_table(
        table_before_paragraph(goal_caption),
        ["Mục tiêu", "Bằng chứng hiện tại", "Đánh giá", "Kết luận học thuật", "Việc còn lại"],
        [
            ["Web App BookVerse AI", "Typecheck, lint, 258 unit test và production build đạt", "Đạt phần kỹ thuật lõi", "Có thể demo trên dữ liệu cục bộ", "Sửa 15/92 E2E đang lỗi"],
            ["Recommendation", "Behavior HitRate@10 = 0,021008; Hybrid = 0,008403", "Chưa đạt ưu thế", "NO_PROMOTION; giữ cấu hình hiện tại", "Thu dữ liệu thật và tạo final_v2 chưa quan sát"],
            ["RAG Assistant", "Unit test grounding/citation đạt; chưa có human relevance chính thức", "Đạt contract kỹ thuật", "Không công bố accuracy người dùng", "Thực hiện blinded relevance study"],
            ["Dataset", "2.200 sách synthetic + 3.046 metadata Open Library", "Đạt pipeline/provenance", "Không đại diện hành vi production", "UAT/telemetry có consent"],
            ["Ebook, P2P, Admin", "Luồng chức năng và dữ liệu demo sẵn sàng", "Đạt phạm vi demo", "Thanh toán chỉ là Sandbox", "Tích hợp cổng thật ngoài phạm vi đồ án"],
        ],
        [3.5, 4.8, 3.0, 4.4, 4.3],
    )
    replace_paragraph(goal_caption,
                      "Bảng 5-1: Đối chiếu mục tiêu với bằng chứng kiểm thử hiện hành",
                      size=11.5, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                      indent=False, before=4, after=8)

    replace_paragraph(
        find_paragraph(doc, "5.2.1 Quy mô dữ liệu thử nghiệm: Dữ liệu thử nghiệm hiện tại gồm hơn 150 đầu sách chuẩn và hàng trăm người dùng mô phỏng; cần mở rộng kiểm thử trên kho dữ liệu hàng chục ngàn sách thực tế."),
        "5.2.1 Quy mô và nguồn dữ liệu: Benchmark dùng 2.200 sách cùng dữ liệu hành vi synthetic; catalog tuyển chọn gồm 3.046 metadata Open Library nhưng không có interaction thật. Do đó kết quả offline chỉ chứng minh pipeline và tính tái lập, chưa chứng minh hiệu quả production.",
        indent=False,
    )
    replace_paragraph(
        find_paragraph(doc, "Hệ thống mang lại trải nghiệm tích cực cho độc giả thông qua tính năng gợi ý cá nhân hóa hiệu quả, công cụ đọc Ebook trực tuyến tương tác cao, sàn giao dịch sách cũ P2P bảo vệ môi trường, cùng Trợ lý AI Nova hỗ trợ tư vấn 24/7. Thử nghiệm thực nghiệm chứng minh hệ thống hoạt động ổn định, tốc độ phản hồi nhanh và độ chính xác gợi ý cải thiện rõ rệt."),
        "Hệ thống đã hoàn thiện các luồng trình diễn chính gồm catalog, đọc Ebook, hội viên, P2P marketplace, thống kê đọc và trợ lý RAG có fallback. Temporal evaluation cho thấy Behavior baseline đang mạnh hơn Hybrid production; vì vậy đồ án không tuyên bố mô hình lai đã cải thiện độ chính xác. Đóng góp chính nằm ở pipeline đánh giá chống leakage, provenance guard và khả năng tái lập artifact.",
    )

    replace_paragraph(
        find_paragraph(doc, "Trả lời: pgvector lưu trữ các vector embedding 1536 chiều của sách và các đoạn văn bản (chunks). Khi người dùng hỏi Trợ lý AI Nova, hệ thống thực hiện tìm kiếm vector tương đồng cosine (<=>) với chỉ mục HNSW (m=16, ef_construction=64) để tìm nhanh các đoạn sách liên quan nhất dưới 50ms và cung cấp context chính xác cho LLM."),
        "Trả lời: pgvector cung cấp kiểu dữ liệu vector và toán tử khoảng cách cosine <=>. Schema hiện dùng kiểu vector chưa khóa dimension; migration chưa tạo HNSW/IVFFlat nên truy vấn hiện tại không được mô tả là ANN dưới 50 ms. Theo tài liệu pgvector, HNSW có thể cải thiện speed-recall trade-off nhưng tốn thời gian build và bộ nhớ; đây là hướng tối ưu sau khi khóa embedding model, dimension và benchmark trên dữ liệu đủ lớn [7].",
        indent=False,
    )

    # Bổ sung tài liệu tham khảo chính thức cho hai nguồn dữ liệu và temporal split.
    reference_anchor = find_paragraph(doc, "PHỤ LỤC A — DANH SÁCH TỪ VIẾT TẮT")
    references = [
        "[10] Google, 'Using the Google Books API,' Google for Developers. [Online]. Available: https://developers.google.com/books/docs/v1/using. Accessed: Aug. 31, 2026.",
        "[11] Internet Archive, 'Open Library APIs and Usage Guidelines.' [Online]. Available: https://openlibrary.org/developers/api. Accessed: Aug. 31, 2026.",
        "[12] Y. Ji, A. Sun, J. Zhang, and C. Li, 'A Critical Study on Data Leakage in Recommender System Offline Evaluation,' ACM Trans. Inf. Syst., vol. 41, no. 3, pp. 75:1-75:27, 2023, doi: 10.1145/3569930.",
        "[13] pgvector contributors, 'pgvector: Open-source vector similarity search for PostgreSQL.' [Online]. Available: https://github.com/pgvector/pgvector. Accessed: Aug. 31, 2026.",
    ]
    for reference in references:
        insert_paragraph(reference_anchor, reference)


def update_checklist(doc) -> None:
    # Tìm theo tiêu đề thay vì dùng chỉ số cố định: sau khi chèn
    # các bảng ở Mục 2.2, vị trí của bảng checklist sẽ thay đổi.
    table = next(
        table
        for table in doc.tables
        if table.rows
        and "Nội dung góp ý của Hội đồng" in table.cell(0, 0).text
    )
    values = [
        "Quy trình thu thập dataset & Pipeline",
        "Mục 2.2 mô tả hai nguồn API, quy tắc thu thập, chuẩn hóa, quality gate, provenance, temporal split và giới hạn sử dụng.",
        "ĐÃ CẬP NHẬT",
    ]
    matching_row = next(
        (
            row
            for row in table.rows[1:]
            if "quy trình thu thập dataset" in row.cells[0].text.lower()
        ),
        None,
    )
    row = matching_row.cells if matching_row is not None else table.add_row().cells
    for index, value in enumerate(values):
        set_cell_text(row[index], value, align=WD_ALIGN_PARAGRAPH.LEFT)


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)
    build_pipeline_image()
    build_evaluation_image()
    document = Document(str(SOURCE_DOCX))
    insert_dataset_section(document)
    revise_existing_content(document)
    update_checklist(document)

    # Yêu cầu Word cập nhật các field (PAGE/TOC) khi mở tệp.
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(str(OUTPUT_DOCX))
    print(f"Đã tạo: {OUTPUT_DOCX}")
    print(f"Kích thước: {OUTPUT_DOCX.stat().st_size} bytes")


if __name__ == "__main__":
    main()
