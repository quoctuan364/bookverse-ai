# -*- coding: utf-8 -*-
"""Cập nhật luận văn bằng số liệu benchmark mới và quy trình dữ liệu người dùng thật.

Script chỉ đọc bản luận văn đã cập nhật trước đó và luôn ghi sang tệp mới.
Không xóa hoặc ghi đè bất kỳ tệp dữ liệu/luận văn nguồn nào.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from revise_thesis_dataset_feedback import (
    ASSET_DIR,
    fill_table,
    find_paragraph,
    find_paragraph_contains,
    insert_paragraph,
    insert_table,
    pil_font,
    replace_paragraph,
    set_cell_text,
    style_paragraph,
    table_before_paragraph,
)


ROOT = Path(r"D:\Doantotnghiep")
SOURCE_DOCX = ROOT / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_CAP_NHAT_GOP_Y_DATASET.docx"
OUTPUT_DOCX = ROOT / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_CAP_NHAT_DU_LIEU_THUC_HYBRID.docx"
EVALUATION_IMAGE = ASSET_DIR / "evaluation_hybrid_latest.png"


def build_latest_evaluation_image() -> None:
    """Tạo biểu đồ NDCG@10 và HitRate@10 từ artifact benchmark mới nhất."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    methods = ["Popularity", "Content", "Behavior", "Hybrid\nFixed", "Hybrid\nDynamic", "Hybrid\nRRF", "Random\nSanity"]
    ndcg = [0.003395, 0.001876, 0.002928, 0.001865, 0.001865, 0.003088, 0.003601]
    hit_rate = [0.013830, 0.007447, 0.013830, 0.007447, 0.007447, 0.015957, 0.017021]

    image = Image.new("RGB", (1900, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (950, 55),
        "TEMPORAL EVALUATION TRÊN DỮ LIỆU SYNTHETIC TẠI K = 10",
        font=pil_font(42, bold=True),
        fill="#173F5F",
        anchor="mm",
    )
    x1, y1, x2, y2 = 135, 150, 1810, 775
    max_value = 0.020
    for tick in range(0, 21, 4):
        value = tick / 1000
        y = y2 - (value / max_value) * (y2 - y1)
        draw.line((x1, y, x2, y), fill="#D7DEE2", width=2)
        draw.text((x1 - 18, y), f"{value:.3f}", font=pil_font(20), fill="#56636A", anchor="rm")
    draw.line((x1, y1, x1, y2), fill="#34495E", width=4)
    draw.line((x1, y2, x2, y2), fill="#34495E", width=4)

    group_width = (x2 - x1) / len(methods)
    bar_width = 48
    for index, method in enumerate(methods):
        center = x1 + group_width * (index + 0.5)
        for offset, value, color in [(-30, ndcg[index], "#2A7F78"), (30, hit_rate[index], "#D77A61")]:
            left = center + offset - bar_width / 2
            right = center + offset + bar_width / 2
            top = y2 - (value / max_value) * (y2 - y1)
            draw.rounded_rectangle((left, top, right, y2), radius=7, fill=color)
            draw.text(((left + right) / 2, top - 9), f"{value:.4f}", font=pil_font(17), fill="#263238", anchor="ms")
        draw.multiline_text(
            (center, y2 + 30), method, font=pil_font(20, bold=True),
            fill="#263238", anchor="ma", align="center", spacing=3,
        )

    draw.rounded_rectangle((610, 895, 655, 932), radius=6, fill="#2A7F78")
    draw.text((675, 913), "NDCG@10", font=pil_font(23), fill="#263238", anchor="lm")
    draw.rounded_rectangle((1000, 895, 1045, 932), radius=6, fill="#D77A61")
    draw.text((1065, 913), "HitRate@10", font=pil_font(23), fill="#263238", anchor="lm")
    image.save(EVALUATION_IMAGE, quality=95)


def replace_by_prefix(document, prefix: str, text: str, **style):
    """Thay đoạn theo tiền tố để script bền hơn khi nội dung dài thay đổi nhẹ."""
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return replace_paragraph(paragraph, text, **style)
    raise ValueError(f"Không tìm thấy đoạn bắt đầu bằng: {prefix}")


def update_dataset_section(document) -> None:
    replace_by_prefix(
        document,
        "Benchmark chỉ dùng lớp synthetic 2.200 sách",
        "Benchmark dùng artifact synthetic gồm 2.200 sách và 18.000 InteractionEvent; 2.000 sách ACTIVE được đưa vào candidate snapshot. Strong-positive gồm purchase hợp lệ, ReadingSession từ 300 giây hoặc progressPercent từ 50%, Bookmark, Favorite và Review từ 4 sao. Các tín hiệu yếu như VIEW, SEARCH và IMPRESSION không được xem là ground truth; riêng IMPRESSION có trọng số huấn luyện bằng 0.",
    )
    replace_by_prefix(
        document,
        "Mốc global temporal cutoff là 01/06/2026",
        "Đánh giá dùng hai mốc thời gian cố định: train có 9.959 positive trước 01/05/2026; validation có 2.407 positive trong khoảng [01/05/2026, 01/06/2026) với 906 người dùng đủ điều kiện; test có 2.584 positive từ 01/06/2026 với 940 người dùng đủ điều kiện. Cách chia train-validation-test theo thời gian hạn chế data leakage và cho phép chọn cấu hình trên validation trước khi báo cáo test [12].",
    )

    benchmark_caption = find_paragraph(document, "Bảng 2-3: Thành phần dataset dùng cho temporal evaluation")
    fill_table(
        table_before_paragraph(benchmark_caption),
        ["Thành phần benchmark", "Số lượng / Giá trị"],
        [
            ["Book nguồn / candidate", "2.200 / 2.000"],
            ["InteractionEvent / ReadingSession", "18.000 / 6.200"],
            ["Bookmark / Review / OrderItem", "2.800 / 3.600 / 4.708"],
            ["Positive sau policy", "14.950"],
            ["Train / Validation / Test positive", "9.959 / 2.407 / 2.584"],
            ["User validation / test đủ điều kiện", "906 / 940"],
            ["Fingerprint dataset", "a19c20e3...02963"],
            ["Checksum kết quả", "f8368bd3...d69928"],
        ],
        [9.5, 8.5],
    )

    old_limit_heading = find_paragraph(document, "2.2.5. Giới hạn và nguyên tắc diễn giải")
    replace_paragraph(
        old_limit_heading,
        "2.2.6. Giới hạn và nguyên tắc diễn giải",
        bold=True,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        indent=False,
        before=7,
        after=3,
    )

    insert_paragraph(old_limit_heading, "2.2.5. Quy trình thu thập tương tác người dùng thật", kind="h2")
    insert_paragraph(
        old_limit_heading,
        "Hệ thống đã bổ sung hạ tầng telemetry để thu dữ liệu thật có kiểm soát. Đây là cơ chế sẵn sàng cho pilot, không phải bằng chứng rằng dữ liệu người dùng thật đã được thu. Tại thời điểm khóa luận, REAL_USER_DATA = 0 và trạng thái đánh giá online là BLOCKED_BY_DATA.",
    )
    steps = [
        "Người dùng chủ động đồng ý tham gia nghiên cứu (opt-in consent). Nếu chưa đồng ý hoặc đã rút đồng ý, hệ thống không ghi sự kiện nghiên cứu.",
        "Giao diện phát sinh 9 loại sự kiện: IMPRESSION, RECOMMENDATION_CLICK, VIEW, SEARCH, FAVORITE, BOOKMARK, ADD_TO_CART, PURCHASE và RATING. IMPRESSION chỉ mô tả lần hiển thị, không được xem là phản hồi tích cực.",
        "Dịch vụ phía server xác thực user, book, kiểu sự kiện và metadata; khóa idempotency loại bản ghi trùng do retry hoặc double-click trước khi ghi database.",
        "Mỗi bản ghi lưu thời điểm, session, vị trí hiển thị, nguồn gợi ý/model version và provenance. Dữ liệu cá nhân không cần thiết không được đưa vào file benchmark.",
        "Trang quản trị /admin/research-data chỉ tổng hợp dữ liệu đã consent; khi xuất nghiên cứu, định danh người dùng phải được giả danh hóa và kèm manifest/checksum để tái lập.",
        "Chỉ mở đánh giá pilot sau khi đạt đồng thời ít nhất 30 người dùng đã consent, 500 impression, 50 outcome và thời gian quan sát 28 ngày. Sau đó tạo temporal split mới và benchmark lại trước khi thay đổi mô hình production.",
    ]
    for step in steps:
        insert_paragraph(old_limit_heading, step, kind="number")

    insert_table(
        document,
        old_limit_heading,
        ["Nhóm tín hiệu", "Sự kiện", "Vai trò trong đánh giá"],
        [
            ["Exposure", "IMPRESSION", "Mẫu số cho CTR; trọng số học = 0"],
            ["Quan tâm", "RECOMMENDATION_CLICK, VIEW, SEARCH", "Tín hiệu ý định yếu; không tự động tạo ground truth"],
            ["Ý định mạnh", "FAVORITE, BOOKMARK, ADD_TO_CART", "Tín hiệu tích cực có điều kiện và cần chống trùng"],
            ["Kết quả", "PURCHASE, RATING", "Outcome mạnh; phải kiểm tra đơn hợp lệ và thang điểm"],
        ],
        [3.0, 7.8, 7.2],
    )
    insert_paragraph(old_limit_heading, "Bảng 2-4: Phân loại sự kiện telemetry phục vụ đánh giá người dùng thật", kind="caption")

    technology_caption = find_paragraph(document, "Bảng 2-4: Bảng công nghệ sử dụng trong BookVerse AI")
    replace_paragraph(
        technology_caption,
        "Bảng 2-5: Bảng công nghệ sử dụng trong BookVerse AI",
        size=11.5,
        bold=True,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False,
        before=4,
        after=8,
    )

    replace_by_prefix(
        document,
        "18.000 sự kiện là synthetic",
        "18.000 sự kiện hiện tại là synthetic, không đại diện hành vi người dùng thật; số liệu telemetry thật hiện bằng 0 nên CTR, UAT và SUS vẫn là NOT_AVAILABLE.",
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        indent=False,
        after=3,
    )
    replace_by_prefix(
        document,
        "Khi thu được dữ liệu người dùng thật có consent",
        "Hạ tầng consent và telemetry đã sẵn sàng nhưng chưa đạt data gate 30 người dùng / 500 impression / 50 outcome / 28 ngày; mọi kết luận production phải chờ pilot và benchmark lại.",
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        indent=False,
        after=3,
    )

    toc_old = find_paragraph_contains(document, "2.2.5. Giới hạn và nguyên tắc diễn giải")
    replace_paragraph(
        toc_old,
        "    2.2.5. Quy trình thu thập tương tác người dùng thật ............... 13",
        size=12,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        indent=False,
        before=0,
        after=0,
        line_spacing=1.0,
    )
    toc_anchor = find_paragraph_contains(document, "2.3. Công nghệ sử dụng trong BookVerse AI")
    toc_line = toc_anchor.insert_paragraph_before("    2.2.6. Giới hạn và nguyên tắc diễn giải .............................. 15")
    style_paragraph(toc_line, size=12, alignment=WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=0, after=0, line_spacing=1.0)


def update_evaluation_section(document) -> None:
    caption = find_paragraph(document, "Bảng 4-1: Kết quả temporal evaluation trên cùng cohort tại K = 10")
    rows = [
        ["Popularity", "0,001383", "0,006223", "0,013830", "0,003395", "0,003897", "0,0065"],
        ["Content", "0,000745", "0,002837", "0,007447", "0,001876", "0,002402", "0,8180"],
        ["Behavior", "0,001383", "0,004495", "0,013830", "0,002928", "0,004530", "0,9660"],
        ["Hybrid Fixed", "0,000745", "0,002748", "0,007447", "0,001865", "0,002501", "0,8125"],
        ["Hybrid Dynamic", "0,000745", "0,002748", "0,007447", "0,001865", "0,002501", "0,8125"],
        ["Hybrid RRF", "0,001596", "0,005612", "0,015957", "0,003088", "0,003689", "0,9565"],
        ["Random sanity", "0,001702", "0,005205", "0,017021", "0,003601", "0,005686", "0,9895"],
    ]
    fill_table(
        table_before_paragraph(caption),
        ["Phương pháp", "P@10", "R@10", "HR@10", "NDCG@10", "MRR@10", "Coverage"],
        rows,
        [3.4, 2.1, 2.1, 2.1, 2.3, 2.1, 2.4],
    )

    figure_caption = find_paragraph_contains(document, "Hình 4-1: NDCG@10 và HitRate@10")
    previous = figure_caption._p.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        for child in list(previous):
            if child.tag != qn("w:pPr"):
                previous.remove(child)
        picture_paragraph = Paragraph(previous, figure_caption._parent)
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.add_run().add_picture(str(EVALUATION_IMAGE), width=Cm(17.0))
    replace_paragraph(
        figure_caption,
        "Hình 4-1: NDCG@10 và HitRate@10 của bảy phương pháp trên test synthetic",
        size=11.5,
        bold=True,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False,
        before=4,
        after=8,
    )

    anchor = find_paragraph(document, "4.3. Xử lý các trường hợp ngoại lệ")
    insert_paragraph(
        anchor,
        "Artifact benchmark có fingerprint a19c20e3...02963 và checksum f8368bd3...d69928. Cả ba cấu hình được thử trên validation đều đạt NDCG@10 = 0,002591 và Recall@10 = 0,003679; do Dynamic Alpha không cải thiện so với cấu hình cố định, hệ thống giữ feature flag HYBRID_DYNAMIC_ALPHA = false và quyết định NO_PROMOTION.",
    )
    insert_paragraph(
        anchor,
        "Trên test, Hybrid RRF là biến thể lai tốt nhất theo HitRate@10 (0,015957), nhưng NDCG@10 (0,003088) vẫn thấp hơn Popularity (0,003395) và Random seeded sanity (0,003601). Random đứng đầu một số chỉ số cho thấy dữ liệu synthetic chưa biểu diễn đủ cấu trúc sở thích có thể học; đây là tín hiệu cần cải thiện dữ liệu và protocol, không phải bằng chứng Random phù hợp cho production.",
    )
    insert_paragraph(
        anchor,
        "Toàn bộ 940 người dùng test thuộc cohort warm; cold = 0 và sparse = 0. Vì vậy benchmark này chưa thể chứng minh Dynamic Alpha xử lý cold-start. Coverage cao chỉ cho biết danh sách gợi ý trải rộng trên catalog, không đồng nghĩa độ liên quan, mức hài lòng hoặc giảm filter bubble.",
    )


def update_conclusion_and_defense(document) -> None:
    goal_caption = find_paragraph(document, "Bảng 5-1: Đối chiếu mục tiêu với bằng chứng kiểm thử hiện hành")
    goal_table = table_before_paragraph(goal_caption)
    for row in goal_table.rows[1:]:
        label = row.cells[0].text.strip()
        if label == "Web App BookVerse AI":
            values = [
                "Web App BookVerse AI",
                "263/263 TypeScript unit; 82/82 Python; integration telemetry 7/7; typecheck và lint đạt",
                "Đạt phần kỹ thuật lõi",
                "Có thể demo trên dữ liệu cục bộ",
                "Chạy lại production build sau lần kiểm tra bị dừng giữa chừng",
            ]
        elif label == "Recommendation":
            values = [
                "Recommendation",
                "RRF HR@10 = 0,015957; Fixed/Dynamic = 0,007447; Random NDCG@10 = 0,003601",
                "Hybrid chưa đạt ưu thế ổn định",
                "NO_PROMOTION; Dynamic Alpha tắt",
                "Thu dữ liệu thật, bổ sung cold/sparse cohort và benchmark lại",
            ]
        elif label == "Dataset":
            values = [
                "Dataset",
                "2.200 sách + 18.000 event synthetic; 3.046 metadata Open Library; telemetry thật = 0",
                "Đạt pipeline/provenance",
                "Chưa đại diện hành vi production",
                "Pilot ≥30 user, 500 impression, 50 outcome, 28 ngày",
            ]
        else:
            continue
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value, align=WD_ALIGN_PARAGRAPH.LEFT)

    replace_by_prefix(
        document,
        "5.2.1 Quy mô và nguồn dữ liệu:",
        "5.2.1 Quy mô và nguồn dữ liệu: Benchmark dùng 2.200 sách và 18.000 sự kiện synthetic; catalog tuyển chọn gồm 3.046 metadata Open Library. Hạ tầng consent/telemetry đã được triển khai và kiểm thử, nhưng hiện chưa có interaction từ người dùng thật (REAL_USER_DATA = 0). Do đó kết quả offline chỉ chứng minh pipeline và tính tái lập, chưa chứng minh hiệu quả production.",
        indent=False,
    )
    replace_by_prefix(
        document,
        "Hệ thống đã hoàn thiện các luồng trình diễn chính",
        "Hệ thống đã hoàn thiện các luồng trình diễn chính và hạ tầng thu 9 loại sự kiện có consent, idempotency và provenance. Benchmark temporal mới so sánh bảy phương pháp cho thấy Hybrid RRF cải thiện HitRate so với Hybrid Fixed/Dynamic nhưng chưa vượt ổn định các baseline; Dynamic Alpha không được bật. Đóng góp chính của đồ án nằm ở pipeline đánh giá chống leakage, cơ chế thu dữ liệu thật có kiểm soát và khả năng tái lập artifact; hiệu quả với người dùng thật là công việc tiếp theo.",
    )

    replace_by_prefix(
        document,
        "Trả lời: Đồ án sử dụng kiến trúc dữ liệu 2 lớp minh bạch:",
        "Trả lời: Đồ án tách ba phạm vi dữ liệu: (1) 2.200 sách và 18.000 tương tác synthetic dùng để kiểm thử pipeline/temporal evaluation; (2) 3.046 metadata Open Library dùng cho catalog; (3) hạ tầng telemetry thật đã sẵn sàng nhưng hiện REAL_USER_DATA = 0. Vì vậy đồ án không tuyên bố dữ liệu synthetic đại diện hành vi production.",
        indent=False,
    )
    replace_by_prefix(
        document,
        "Trả lời: Qua đánh giá thực nghiệm trung thực",
        "Trả lời: Benchmark mới cho thấy Hybrid Fixed và Dynamic có NDCG@10 = 0,001865; Hybrid RRF đạt 0,003088 nhưng vẫn thấp hơn Popularity 0,003395 và Random sanity 0,003601. Dynamic Alpha không tạo khác biệt vì toàn bộ 940 user test đều thuộc cohort warm. Kết quả này dẫn tới quyết định NO_PROMOTION và yêu cầu thu dữ liệu thật/cold-sparse cohort trước khi tối ưu tiếp.",
        indent=False,
    )
    replace_by_prefix(
        document,
        "Trả lời: Trong Recommender System, Random Split",
        "Trả lời: Random split có nguy cơ đưa tương tác tương lai vào train. Đề tài dùng train trước 01/05/2026 (9.959 positive), validation trong tháng 05/2026 (2.407 positive) để chọn cấu hình, và test từ 01/06/2026 (2.584 positive) để báo cáo cuối. Candidate snapshot được khóa theo thời gian nhằm giảm data leakage.",
        indent=False,
    )
    replace_by_prefix(
        document,
        "Trả lời: (1) Thu thập tập dữ liệu tương tác từ người dùng thật",
        "Trả lời: (1) Chạy pilot tối thiểu 30 người dùng đã consent trong 28 ngày, đạt 500 impression và 50 outcome; (2) xuất dữ liệu giả danh hóa kèm manifest/checksum; (3) tạo temporal split mới có cold/sparse/warm cohort; (4) so sánh lại Fixed, Dynamic Alpha, RRF và baseline; (5) chỉ bật mô hình mới khi vượt tiêu chí đã định và có UAT/SUS chính thức.",
        indent=False,
    )

    checklist = next(
        table for table in document.tables
        if table.rows and "Nội dung góp ý của Hội đồng" in table.cell(0, 0).text
    )
    for row in checklist.rows[1:]:
        item = row.cells[0].text.lower()
        if "quy trình thu thập dataset" in item:
            set_cell_text(row.cells[1], "Mục 2.2 bổ sung nguồn metadata/synthetic, telemetry 9 sự kiện, consent, idempotency, provenance và data gate cho pilot người dùng thật.")
            set_cell_text(row.cells[2], "ĐÃ CẬP NHẬT")
        elif "minh chứng chỉ số" in item:
            set_cell_text(row.cells[1], "Bảng 4-1 cập nhật 7 phương pháp, 6 metric tại K=10, validation tuning, cohort, fingerprint/checksum và quyết định NO_PROMOTION.")
            set_cell_text(row.cells[2], "ĐÃ CẬP NHẬT")


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)
    build_latest_evaluation_image()
    document = Document(str(SOURCE_DOCX))
    update_dataset_section(document)
    update_evaluation_section(document)
    update_conclusion_and_defense(document)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(str(OUTPUT_DOCX))
    print(f"Đã tạo: {OUTPUT_DOCX}")
    print(f"Kích thước: {OUTPUT_DOCX.stat().st_size} bytes")


if __name__ == "__main__":
    main()
