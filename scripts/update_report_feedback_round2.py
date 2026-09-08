from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


WORKSPACE = Path(r"D:\Doantotnghiep")
SOURCE = WORKSPACE / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098.docx"
OUTPUT = WORKSPACE / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_HOAN_THIEN_GOP_Y_LAN_2.docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Đổi nội dung nhưng giữ thuộc tính đoạn và định dạng run đầu tiên."""
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_rpr = deepcopy(paragraph.runs[0]._r.rPr)

    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._r.insert(0, first_rpr)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    """Chèn đoạn ngay sau đoạn hiện tại."""
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = Paragraph(new_p, paragraph._parent)
    if style:
        result.style = style
    result.add_run(text)
    return result


def find_paragraph(document: Document, startswith: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"Cần đúng 1 đoạn bắt đầu bằng {startswith!r}, tìm thấy {len(matches)}")
    return matches[0]


def find_heading(document: Document, startswith: str, level: int = 1) -> Paragraph:
    style_name = f"Heading {level}"
    matches = [
        p
        for p in document.paragraphs
        if p.style.name == style_name and p.text.strip().startswith(startswith)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Cần đúng 1 heading {level} bắt đầu bằng {startswith!r}, tìm thấy {len(matches)}"
        )
    return matches[0]


def replace_in_all_paragraphs(document: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace(old, new))
            count += 1
    return count


def add_before(anchor: Paragraph, text: str, style: str) -> Paragraph:
    return anchor.insert_paragraph_before(text, style=style)


def main() -> None:
    parser = argparse.ArgumentParser(description="Hoàn thiện báo cáo theo góp ý lần 2")
    parser.add_argument(
        "--metadata-only",
        action="store_true",
        help="Chỉ chuẩn hóa metadata của file đầu ra sau khi Word cập nhật field",
    )
    args = parser.parse_args()

    if args.metadata_only:
        if not OUTPUT.exists():
            raise FileNotFoundError(OUTPUT)
        final_doc = Document(OUTPUT)
        final_doc.core_properties.author = "Lương Nguyễn Quốc Tuấn"
        final_doc.core_properties.last_modified_by = "Lương Nguyễn Quốc Tuấn"
        final_doc.core_properties.subject = "Đồ án tốt nghiệp BookVerse AI – bản hoàn thiện theo góp ý lần 2"
        final_doc.save(OUTPUT)
        print(f"Đã chuẩn hóa metadata: {OUTPUT}")
        return

    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(SOURCE)
    doc.core_properties.author = "Lương Nguyễn Quốc Tuấn"
    doc.core_properties.last_modified_by = "Lương Nguyễn Quốc Tuấn"
    doc.core_properties.subject = "Đồ án tốt nghiệp BookVerse AI – bản hoàn thiện theo góp ý lần 2"

    # 1) Chuẩn hóa phạm vi người dùng: Guest, Reader, Seller, Admin.
    scope_cell = doc.tables[2].cell(4, 0)
    for paragraph in scope_cell.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Phạm vi người dùng:"):
            set_paragraph_text(
                paragraph,
                "Phạm vi người dùng: khách truy cập, độc giả đã đăng ký, người bán và quản trị viên. "
                "Người bán là độc giả được cấp hoặc nâng quyền để đăng tin, quản lý tin bán và xử lý giao dịch sách cũ P2P. "
                "Hệ thống tập trung vào quy trình khám phá sách, mô phỏng mua sách, đăng ký gói hội viên, "
                "đọc thử tối đa 10% khi chưa có quyền, đọc toàn bộ nội dung minh họa khi có quyền truy cập, "
                "tương tác với sách và quản trị dữ liệu.",
            )
        elif text.startswith("Đối tượng phục vụ gồm"):
            set_paragraph_text(
                paragraph,
                "Đối tượng phục vụ gồm độc giả cần tìm, mua và đọc sách; người bán là độc giả được nâng quyền để tham gia "
                "giao dịch sách cũ P2P; người dùng muốn được tư vấn theo sở thích; và quản trị viên chịu trách nhiệm quản lý "
                "dữ liệu sách, tài khoản, đơn hàng và báo cáo hệ thống.",
            )

    actor_table = doc.tables[10]
    nova_rows = [row for row in actor_table.rows if row.cells[0].text.strip() == "Trợ lý AI Nova"]
    if len(nova_rows) != 1:
        raise RuntimeError(f"Không xác định duy nhất dòng Nova trong bảng actor: {len(nova_rows)}")
    nova_row = nova_rows[0]
    nova_row._tr.getparent().remove(nova_row._tr)

    # 2) Sửa phần khảo sát hệ thống tương tự, bỏ lặp chữ và tránh nhận định tuyệt đối.
    set_paragraph_text(
        find_paragraph(doc, "Khảo sát các nền tảng thương mại điện tử sách hiện nay"),
        "Khảo sát đối chiếu một số nền tảng sách và nội dung số hiện nay (Fahasa, Tiki, Goodreads và Waka):",
    )
    set_paragraph_text(
        find_paragraph(doc, "• Fahasa & Tiki:"),
        "• Fahasa và Tiki: Hai nền tảng hỗ trợ catalog và mua sách trực tuyến. Báo cáo chỉ đối chiếu các chức năng "
        "công khai; không kết luận về thuật toán gợi ý nội bộ do không có tài liệu kỹ thuật để kiểm chứng.",
    )
    set_paragraph_text(
        find_paragraph(doc, "• Goodreads:"),
        "• Goodreads: Nền tảng cộng đồng sách hỗ trợ đánh giá và theo dõi hoạt động đọc. BookVerse AI khác về phạm vi "
        "prototype khi tích hợp mua sách, gói đọc nội dung số và sàn sách cũ P2P trong cùng hệ thống.",
    )
    set_paragraph_text(
        find_paragraph(doc, "• Waka:"),
        "• Waka: Nền tảng nội dung số và Ebook. Báo cáo chỉ đối chiếu nhóm chức năng được công khai, không đánh giá "
        "tuyệt đối chất lượng AI hoặc kiến trúc nội bộ của nền tảng.",
    )

    # 3) Giảm mức khẳng định đối với RAG và Hybrid theo đúng bằng chứng hiện có.
    set_paragraph_text(
        find_paragraph(doc, "• 3. Trợ lý AI Nova RAG:"),
        "• 3. Trợ lý AI Nova RAG: Thành phần RAG hỗ trợ truy xuất ngữ cảnh, sinh câu trả lời có dẫn nguồn khi đủ bằng chứng "
        "và chuyển sang fallback an toàn khi thiếu dữ liệu. Đồ án chưa có tập đánh giá do người chấm gán nhãn, vì vậy không "
        "công bố độ chính xác câu trả lời của RAG.",
    )
    hybrid_goal_cell = doc.tables[4].cell(3, 2)
    set_paragraph_text(
        hybrid_goal_cell.paragraphs[0],
        "So sánh các cấu hình trên cùng temporal split; chỉ chọn cấu hình triển khai khi vượt tiêu chí đã khóa. "
        "Kết quả hiện tại là NO_PROMOTION và Dynamic Alpha đang tắt.",
    )

    # 4) Mô tả chính xác embedding theo mã nguồn và giới hạn đo lường.
    embedding_p = find_paragraph(doc, "Schema BookVerse AI hỗ trợ lưu vector embedding")
    set_paragraph_text(
        embedding_p,
        "Trong lib/book-embeddings.ts, provider mặc định là OpenAI với model text-embedding-3-small; mã không truyền tham số "
        "dimensions nên dùng kích thước native 1.536 chiều và kiểm tra độ dài vector từ phản hồi của provider. Khi cấu hình "
        "BOOKVERSE_EMBEDDING_PROVIDER=gemini, hệ thống dùng model text-embedding-004. Chuỗi đầu vào gồm tiêu đề, tác giả, "
        "thể loại, cấp độ (nếu có), tags và mô tả. Vector được lưu bằng pgvector và truy vấn cosine exact với toán tử <=>. "
        "Schema hiện dùng kiểu vector chưa khóa số chiều, chưa tạo HNSW/IVFFlat; tại thời điểm đối chiếu, cơ sở dữ liệu kiểm tra "
        "chưa có bản ghi book_embeddings nên nội dung này là bằng chứng cài đặt hạ tầng, chưa phải kết quả đánh giá thực nghiệm.",
    )

    # 5) Thay lý thuyết và công thức CF bằng đúng item-item implicit trong benchmark.
    cf_theory = find_paragraph(doc, "Collaborative Filtering khai thác mối quan hệ")
    set_paragraph_text(
        cf_theory,
        "Trong benchmark của BookVerse AI, Collaborative Filtering được cài đặt theo hướng item-item trên phản hồi ngầm. "
        "Chỉ các sự kiện strong-positive trong tập train được khử trùng lặp theo cặp user-item rồi dùng để tính co-occurrence; "
        "không sử dụng rating dự đoán kiểu user-user và không đưa sự kiện validation/test vào ma trận huấn luyện.",
    )
    cf_label = find_paragraph(doc, "2. Điểm dự đoán Collaborative Filtering:")
    set_paragraph_text(cf_label, "2. Collaborative Filtering item-item trên strong-positive implicit feedback:")
    cf_formula = find_paragraph(doc, "R_hat(u, i)")
    set_paragraph_text(cf_formula, "cᵢ = |{u : (u, i) ∈ T⁺}|;   cᵢⱼ = |{u : (u, i) ∈ T⁺ và (u, j) ∈ T⁺}|")
    cursor = insert_after(
        cf_formula,
        "sim(i, j) = cᵢⱼ / √(cᵢ × cⱼ)",
        style="Normal",
    )
    cursor = insert_after(
        cursor,
        "score(u, j) = ∑ᵢ∈Iᵤ⁺ sim(i, j)",
        style="Normal",
    )
    insert_after(
        cursor,
        "Trong đó T⁺ là tập tương tác strong-positive thuộc train, Iᵤ⁺ là các sách người dùng u đã tương tác tích cực; "
        "các sách đã xem trong lịch sử được loại khỏi danh sách ứng viên trước khi xếp hạng.",
        style="Normal",
    )

    rag_theory = find_paragraph(doc, "Kiến trúc RAG kết hợp sức mạnh")
    set_paragraph_text(
        rag_theory,
        "Kiến trúc RAG kết hợp mô hình ngôn ngữ với truy xuất thông tin. Khi người dùng đặt câu hỏi, hệ thống truy xuất các "
        "đoạn tri thức liên quan từ dữ liệu sách và chính sách BookVerse AI, sau đó đưa chúng vào ngữ cảnh sinh câu trả lời. "
        "Phản hồi được kiểm tra citation/grounding theo contract; khi thiếu bằng chứng hoặc dịch vụ ngoài lỗi, hệ thống dùng "
        "fallback hoặc từ chối an toàn. Cơ chế này giúp kiểm soát câu trả lời nhưng không đồng nghĩa đã chứng minh độ chính xác "
        "nếu chưa có tập câu hỏi gán nhãn và đánh giá con người.",
    )

    # 6) Sửa số mục Chương 3 từ 3.1 -> 3.2... liên tục.
    chapter_paragraphs = doc.paragraphs
    ch3_heading = find_heading(doc, "CHƯƠNG 3:")
    ch4_heading = find_heading(doc, "CHƯƠNG 4:")
    ch3_start = next(i for i, p in enumerate(chapter_paragraphs) if p._p is ch3_heading._p)
    ch4_start = next(i for i, p in enumerate(chapter_paragraphs) if p._p is ch4_heading._p)
    heading_mapping = (("3.3", "3.2"), ("3.4", "3.3"), ("3.5", "3.4"), ("3.6", "3.5"), ("3.7", "3.6"))
    for paragraph in chapter_paragraphs[ch3_start:ch4_start]:
        if not paragraph.style.name.startswith("Heading"):
            continue
        original = paragraph.text
        for old, new in heading_mapping:
            if re.match(rf"^{re.escape(old)}(?=\.|\s)", original):
                set_paragraph_text(paragraph, re.sub(rf"^{re.escape(old)}", new, original, count=1))
                break

    # 7) Đánh số lại toàn bộ caption hình/bảng Chương 3 theo thứ tự xuất hiện.
    figure_no = 0
    table_no = 0
    for paragraph in chapter_paragraphs[ch3_start:ch4_start]:
        text = paragraph.text.strip()
        fig_match = re.match(r"^Hình(?:\s+3-\d+)?\s*:\s*(.*)$", text, flags=re.IGNORECASE)
        if fig_match:
            figure_no += 1
            set_paragraph_text(paragraph, f"Hình 3-{figure_no}: {fig_match.group(1).strip()}")
            continue
        table_match = re.match(r"^Bảng(?:\s+3-\d+)?\s*:\s*(.*)$", text, flags=re.IGNORECASE)
        if table_match:
            table_no += 1
            set_paragraph_text(paragraph, f"Bảng 3-{table_no}: {table_match.group(1).strip()}")

    # 8) Tái cấu trúc Chương 4 theo trình tự thực nghiệm mà giảng viên yêu cầu.
    h41 = find_heading(doc, "4.1 Các kịch bản thử nghiệm", level=2)
    set_paragraph_text(h41, "4.1. Thiết lập thực nghiệm")
    set_paragraph_text(
        find_paragraph(doc, "Kịch bản 4: Thử nghiệm Trợ lý AI Nova RAG"),
        "Kịch bản 4: Kiểm tra tự động các contract của Trợ lý AI Nova RAG gồm retrieval, citation, paywall và fallback. "
        "Chưa thực hiện đánh giá độ đúng câu trả lời trên tập 30–50 câu hỏi được gán nhãn bởi người chấm.",
    )

    old_h42 = find_heading(doc, "4.2. Kết quả thử nghiệm các kịch bản", level=2)
    add_before(old_h42, "4.2. Dataset và temporal split", "Heading 2")
    add_before(
        old_h42,
        "Benchmark Recommendation sử dụng bộ Ultra Synthetic gồm 2.200 sách và 18.000 sự kiện tương tác. Dữ liệu được chia "
        "theo thời gian thành train/validation/test với lần lượt 9.959/2.407/2.584 sự kiện sau các bước lọc của pipeline; "
        "tham số chỉ được lựa chọn trên validation, còn test được giữ lại để báo cáo. Bộ này phục vụ kiểm tra pipeline, tính "
        "tái lập và benchmark offline, không đại diện cho hành vi người dùng thật.",
        "Normal",
    )
    add_before(old_h42, "4.3. Baseline và chỉ số đánh giá", "Heading 2")
    add_before(
        old_h42,
        "Các phương pháp được so sánh gồm Popularity, Content, Behavior item-item, Hybrid Fixed, Hybrid Dynamic Alpha, "
        "Hybrid RRF và Random seeded sanity check. Tất cả dùng cùng temporal split và K = 10. Các chỉ số báo cáo gồm "
        "Precision@10, Recall@10, HitRate@10, NDCG@10, MRR@10 và Coverage@10; random seed, fingerprint và checksum được "
        "lưu trong artifact để hỗ trợ tái lập.",
        "Normal",
    )
    set_paragraph_text(old_h42, "4.4. Kết quả Recommendation")

    old_h43 = find_heading(doc, "4.3. Xử lý các trường hợp ngoại lệ", level=2)
    add_before(old_h43, "4.5. Đánh giá RAG", "Heading 2")
    add_before(
        old_h43,
        "Repository có các kiểm tra tự động cho retrieval, citation, giới hạn paywall và fallback của RAG. Tuy nhiên, đồ án "
        "chưa xây dựng bộ 30–50 câu hỏi gán nhãn kèm đáp án tham chiếu và chưa có người chấm độc lập; vì vậy các chỉ số "
        "Answer Correctness, Groundedness, Citation Correctness và mức hài lòng người dùng được ghi là NOT_AVAILABLE. "
        "Báo cáo chỉ kết luận cơ chế đã được cài đặt và kiểm tra theo contract, không kết luận độ chính xác của nội dung trả lời.",
        "Normal",
    )
    set_paragraph_text(old_h43, "4.6. Kiểm thử tự động và chức năng")

    ch5 = find_heading(doc, "CHƯƠNG 5.")
    add_before(ch5, "4.7. Hiệu năng", "Heading 2")
    add_before(
        ch5,
        "Đề tài chưa có benchmark độ trễ lặp lại để báo cáo Mean, Median, P95 và Max cho API gợi ý, pgvector hoặc RAG. "
        "Ngưỡng dưới 100 ms, nếu được nêu trong yêu cầu, chỉ là mục tiêu thiết kế cho các truy vấn nội bộ phù hợp; không được "
        "xem là kết quả đã đạt. Kết quả hiệu năng hiện tại: NOT_AVAILABLE.",
        "Normal",
    )
    add_before(ch5, "4.8. Thảo luận và giới hạn thực nghiệm", "Heading 2")
    add_before(
        ch5,
        "Hybrid RRF có HitRate@10 tốt nhất trong nhóm biến thể Hybrid nhưng chưa vượt ổn định mọi baseline; Dynamic Alpha "
        "không cải thiện trên validation nên đang tắt và artifact ghi nhận NO_PROMOTION. Toàn bộ người dùng test hiện thuộc "
        "cohort warm, dữ liệu tương tác là synthetic và chưa có đánh giá RAG bởi người dùng thật. Do đó kết quả mới chứng minh "
        "pipeline có thể chạy và tái lập, chưa đủ để suy rộng hiệu quả production.",
        "Normal",
    )

    # 9) Chương 5: kết luận đúng mức bằng chứng.
    limitation = find_paragraph(doc, "5.2.3 Đàm thoại đa lượt chuyên sâu:")
    set_paragraph_text(
        limitation,
        "5.2.3. Đánh giá RAG và đàm thoại đa lượt: Trợ lý AI Nova đã có kiểm tra contract cho retrieval, citation, paywall "
        "và fallback, nhưng chưa có bộ câu hỏi gán nhãn hoặc đánh giá người dùng để định lượng độ đúng. Khả năng ghi nhớ "
        "ngữ cảnh hội thoại dài qua nhiều phiên cũng chưa được đánh giá thực nghiệm.",
    )

    # Kiểm tra các câu tuyệt đối cũ đã được loại bỏ.
    forbidden = [
        "độ chính xác cao và có trích dẫn",
        "Đạt độ chính xác tối ưu và độ phủ sách rộng khắp",
        "Trợ lý AI Nova hoạt động tốt",
        "Fahasa & Tiki: Fahasa, Tiki:",
        "Goodreads: Goodreads:",
        "Waka: Waka:",
        "R_hat(u, i)",
    ]
    all_text = "\n".join(p.text for p in doc.paragraphs)
    all_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    remaining = [phrase for phrase in forbidden if phrase in all_text]
    if remaining:
        raise RuntimeError(f"Còn nội dung cũ chưa sửa: {remaining}")

    doc.save(OUTPUT)
    print(f"Đã tạo: {OUTPUT}")
    print(f"Caption Chương 3: {figure_no} hình, {table_no} bảng")


if __name__ == "__main__":
    main()
