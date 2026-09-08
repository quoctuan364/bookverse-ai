from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\Doantotnghiep")
SOURCE = ROOT / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_HOAN_THIEN_GOP_Y_LAN_2.docx"
OUTPUT = ROOT / "Do_An_Tot_Nghiep_BookVerse_AI_Luong_Nguyen_Quoc_Tuan_22050098_HOAN_THIEN_GOP_Y_LAN_2_MA_NGUON_THUC_TE_WINDOWS.docx"
FIGURE_DIR = ROOT / "artifacts" / "report_code_figures"

CODE_FONT = Path(r"C:\Windows\Fonts\consola.ttf")
CODE_BOLD_FONT = Path(r"C:\Windows\Fonts\consolab.ttf")
WINDOW_UI_FONT = Path(r"C:\Windows\Fonts\segoeui.ttf")


def read_excerpt(relative_path: str, ranges: list[tuple[int, int]]) -> list[tuple[int | None, str]]:
    """Đọc đúng các dòng mã nguồn; giữa hai đoạn dùng dấu lược để thể hiện trích đoạn."""
    source_path = ROOT / relative_path
    lines = source_path.read_text(encoding="utf-8").splitlines()
    result: list[tuple[int | None, str]] = []
    for index, (start, end) in enumerate(ranges):
        if index:
            result.append((None, "..."))
        for line_number in range(start, end + 1):
            result.append((line_number, lines[line_number - 1].replace("\t", "    ")))
    return result


def draw_highlighted_line(
    draw: ImageDraw.ImageDraw,
    position: tuple[int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
) -> None:
    """Tô màu cú pháp nhẹ để ảnh mã nguồn dễ đọc khi đặt vào Word."""
    x, y = position
    token_pattern = re.compile(
        r'(//.*$|#.*$|"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'|\b\d+(?:\.\d+)?\b|'
        r'\b(?:async|await|const|return|if|else|for|in|def|from|import|function|export|type|interface|'
        r'None|True|False|try|catch|throw|new|continue|class)\b)'
    )
    cursor = 0
    for match in token_pattern.finditer(text):
        prefix = text[cursor : match.start()]
        draw.text((x, y), prefix, font=font, fill="#E6EDF3")
        x += int(draw.textlength(prefix, font=font))

        token = match.group(0)
        if token.startswith(("//", "#")):
            color = "#8B949E"
        elif token.startswith(("\"", "'")):
            color = "#A5D6FF"
        elif token[0].isdigit():
            color = "#FFA657"
        else:
            color = "#FF7B72"
        draw.text((x, y), token, font=font, fill=color)
        x += int(draw.textlength(token, font=font))
        cursor = match.end()

    suffix = text[cursor:]
    draw.text((x, y), suffix, font=font, fill="#E6EDF3")


def render_code_figure(
    relative_path: str,
    ranges: list[tuple[int, int]],
    output_path: Path,
) -> tuple[int, int]:
    rows = read_excerpt(relative_path, ranges)
    width = 2200
    header_height = 108
    line_height = 35
    bottom_padding = 34
    height = header_height + line_height * len(rows) + bottom_padding

    image = Image.new("RGB", (width, height), "#0D1117")
    draw = ImageDraw.Draw(image)
    code_font = ImageFont.truetype(str(CODE_FONT), 25)
    code_bold = ImageFont.truetype(str(CODE_BOLD_FONT), 27)
    window_font = ImageFont.truetype(str(WINDOW_UI_FONT), 28)

    draw.rounded_rectangle((0, 0, width - 1, height - 1), radius=18, outline="#30363D", width=3)
    draw.rectangle((1, 1, width - 2, header_height), fill="#161B22")

    # Thanh tiêu đề theo phong cách VS Code trên Windows, không dùng ba nút macOS.
    draw.rounded_rectangle((28, 27, 82, 81), radius=8, fill="#007ACC")
    draw.text((36, 36), "<>", font=code_bold, fill="#FFFFFF")

    range_label = ", ".join(f"{start}-{end}" for start, end in ranges)
    header = (
        f"VS Code (Windows)  |  {relative_path.replace(chr(92), '/')}  |  "
        f"dòng {range_label}  |  BookVerse AI"
    )
    draw.text((104, 31), header, font=code_bold, fill="#F0F6FC")
    draw.text((width - 205, 31), "—   □   ×", font=window_font, fill="#F0F6FC")

    number_right = 118
    code_x = 150
    y = header_height + 11
    for line_number, code in rows:
        if line_number is None:
            draw.text((code_x, y), code, font=code_bold, fill="#8B949E")
        else:
            number = str(line_number)
            number_width = draw.textlength(number, font=code_font)
            draw.text((number_right - number_width, y), number, font=code_font, fill="#6E7681")
            draw_highlighted_line(draw, (code_x, y), code, code_font)
        y += line_height

    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, format="PNG", optimize=True)
    return image.size


def find_body_caption(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.casefold() == "hình".casefold()
        and paragraph.text.strip().startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Không xác định duy nhất caption {prefix}: {len(matches)}")
    return matches[0]


def replace_caption_text(paragraph, text: str) -> None:
    """Đổi caption nhưng giữ style, căn giữa và định dạng của run đầu tiên."""
    first_run = paragraph.runs[0] if paragraph.runs else None
    bold = first_run.bold if first_run else None
    italic = first_run.italic if first_run else None
    font_name = first_run.font.name if first_run else None
    font_size = first_run.font.size if first_run else None
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font_name
    run.font.size = font_size


def replace_preceding_image(document: Document, caption, image_path: Path, pixel_size: tuple[int, int]) -> str:
    """Thay blob PNG và chỉnh extent theo đúng tỉ lệ để không kéo méo hình."""
    image_paragraph = caption._p.getprevious()
    while image_paragraph is not None and not image_paragraph.xpath(".//a:blip"):
        image_paragraph = image_paragraph.getprevious()
    if image_paragraph is None:
        raise RuntimeError(f"Không tìm thấy hình đứng trước {caption.text}")

    blips = image_paragraph.xpath(".//a:blip")
    if len(blips) != 1:
        raise RuntimeError(f"Số blip không hợp lệ trước {caption.text}: {len(blips)}")
    relationship_id = blips[0].get(qn("r:embed"))
    image_part = document.part.related_parts[relationship_id]
    if image_part.content_type != "image/png":
        raise RuntimeError(f"Ảnh đích không phải PNG: {image_part.content_type}")
    image_part._blob = image_path.read_bytes()

    width_px, height_px = pixel_size
    extents = image_paragraph.xpath(".//wp:extent")
    if len(extents) != 1:
        raise RuntimeError(f"Không xác định được extent của {caption.text}")
    width_emu = int(extents[0].get("cx"))
    height_emu = round(width_emu * height_px / width_px)
    extents[0].set("cy", str(height_emu))
    for shape_extent in image_paragraph.xpath(".//a:xfrm/a:ext"):
        shape_extent.set("cx", str(width_emu))
        shape_extent.set("cy", str(height_emu))

    # Không để ảnh nằm cuối trang còn caption bị đẩy sang trang kế tiếp.
    image_flow_paragraph = Paragraph(image_paragraph, caption._parent)
    image_flow_paragraph.paragraph_format.keep_with_next = True
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = False
    return relationship_id


def main() -> None:
    parser = argparse.ArgumentParser(description="Thay hình mã nguồn cũ trong báo cáo BookVerse AI")
    parser.add_argument(
        "--metadata-only",
        action="store_true",
        help="Chỉ chuẩn hóa tác giả sau khi Word cập nhật field",
    )
    args = parser.parse_args()

    if args.metadata_only:
        if not OUTPUT.exists():
            raise FileNotFoundError(OUTPUT)
        document = Document(OUTPUT)
        document.core_properties.author = "Lương Nguyễn Quốc Tuấn"
        document.core_properties.last_modified_by = "Lương Nguyễn Quốc Tuấn"
        document.core_properties.subject = "Đồ án BookVerse AI - hình mã nguồn khớp repository TypeScript/Python"
        document.save(OUTPUT)
        print(f"Đã chuẩn hóa metadata: {OUTPUT}")
        return

    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    specifications = [
        {
            "prefix": "Hình 3-9:",
            "path": "lib/book-embeddings.ts",
            "ranges": [(39, 60)],
            "caption": "Hình 3-9: Mã tạo nội dung embedding (lib/book-embeddings.ts)",
            "filename": "hinh-3-9-book-embeddings.png",
        },
        {
            "prefix": "Hình 3-10:",
            "path": "lib/assistant-service.ts",
            "ranges": [(318, 349)],
            "caption": "Hình 3-10: Truy vấn cosine pgvector và fallback (lib/assistant-service.ts)",
            "filename": "hinh-3-10-pgvector-cosine.png",
        },
        {
            "prefix": "Hình 3-11:",
            "path": "ai_service/evaluation/hybrid_audit.py",
            "ranges": [(179, 205)],
            "caption": "Hình 3-11: Thuật toán Hybrid RRF (ai_service/evaluation/hybrid_audit.py)",
            "filename": "hinh-3-11-hybrid-rrf.png",
        },
        {
            "prefix": "Hình 3-12:",
            "path": "lib/reader-rag.ts",
            "ranges": [(4, 10), (221, 240)],
            "caption": "Hình 3-12: Prompt grounding và kiểm tra citation (lib/reader-rag.ts)",
            "filename": "hinh-3-12-reader-rag-grounding.png",
        },
    ]

    document = Document(SOURCE)
    replaced: list[str] = []
    for specification in specifications:
        figure_path = FIGURE_DIR / specification["filename"]
        pixel_size = render_code_figure(
            specification["path"], specification["ranges"], figure_path
        )
        caption = find_body_caption(document, specification["prefix"])
        relationship_id = replace_preceding_image(document, caption, figure_path, pixel_size)
        replace_caption_text(caption, specification["caption"])
        replaced.append(f"{specification['prefix']} {relationship_id} -> {figure_path.name}")

    # Các caption UC13-UC29 trước đây dùng style Normal nên bị thiếu trong danh mục.
    # Chuẩn hóa toàn bộ caption Chương 3 để danh mục hình/bảng hiển thị liên tục.
    in_chapter_three = False
    normalized_figures = 0
    normalized_tables = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name == "Heading 1" and text.startswith("CHƯƠNG 3:"):
            in_chapter_three = True
            continue
        if paragraph.style.name == "Heading 1" and text.startswith("CHƯƠNG 4:"):
            break
        if not in_chapter_three:
            continue
        if re.match(r"^Hình 3-\d+:", text):
            paragraph.style = document.styles["HÌnh"]
            normalized_figures += 1
        elif re.match(r"^Bảng 3-\d+:", text):
            paragraph.style = document.styles["Bang"]
            normalized_tables += 1

    document.core_properties.author = "Lương Nguyễn Quốc Tuấn"
    document.core_properties.last_modified_by = "Lương Nguyễn Quốc Tuấn"
    document.core_properties.subject = "Đồ án BookVerse AI - hình mã nguồn khớp repository TypeScript/Python"
    document.save(OUTPUT)

    print(f"Đã tạo: {OUTPUT}")
    for item in replaced:
        print(item)
    print(f"Đã chuẩn hóa style: {normalized_figures} hình, {normalized_tables} bảng Chương 3")


if __name__ == "__main__":
    main()
